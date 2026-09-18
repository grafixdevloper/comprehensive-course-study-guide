#!/usr/bin/env python3
"""
Content extraction script for the Comprehensive Course Study Guide plugin.
Extracts text, structure, and metadata from various file formats with provenance.
"""

import sys
import json
import re
from pathlib import Path
from typing import Dict, List, Any, Optional
from dataclasses import dataclass, asdict
from datetime import datetime

@dataclass
class ExtractedItem:
    source_file: str
    source_type: str  # pdf, pptx, docx, txt, md, image
    page_or_slide: int
    content_type: str  # heading, paragraph, table, list, formula, code, image_caption
    content: str
    metadata: Dict[str, Any]

def extract_pdf(filepath: Path, rel_path: str) -> List[ExtractedItem]:
    """Extract content from PDF with page-level provenance."""
    items = []
    try:
        import fitz
        doc = fitz.open(filepath)

        for page_num in range(doc.page_count):
            page = doc[page_num]

            # Extract text with structure
            blocks = page.get_text("dict")["blocks"]

            for block in blocks:
                if "lines" not in block:
                    continue

                block_text = ""
                for line in block["lines"]:
                    for span in line["spans"]:
                        block_text += span["text"]
                    block_text += "\n"

                block_text = block_text.strip()
                if not block_text:
                    continue

                # Classify content type
                content_type = classify_text_block(block_text, block)

                items.append(ExtractedItem(
                    source_file=rel_path,
                    source_type="pdf",
                    page_or_slide=page_num + 1,
                    content_type=content_type,
                    content=block_text,
                    metadata={
                        "bbox": block.get("bbox", []),
                        "font_info": extract_font_info(block)
                    }
                ))

            # Extract tables
            tables = page.find_tables()
            for table in tables:
                table_data = table.extract()
                if table_data:
                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="pdf",
                        page_or_slide=page_num + 1,
                        content_type="table",
                        content=json.dumps(table_data),
                        metadata={"bbox": table.bbox}
                    ))

            # Extract images info
            images = page.get_images(full=True)
            for img in images:
                items.append(ExtractedItem(
                    source_file=rel_path,
                    source_type="pdf",
                    page_or_slide=page_num + 1,
                    content_type="image",
                    content=f"[Image: {img[0]}]",
                    metadata={"xref": img[0], "width": img[2], "height": img[3]}
                ))

        doc.close()

    except Exception as e:
        print(f"Error extracting PDF {filepath}: {e}", file=sys.stderr)

    return items

def extract_pptx(filepath: Path, rel_path: str) -> List[ExtractedItem]:
    """Extract content from PPTX with slide-level provenance."""
    items = []
    try:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation(filepath)

        for slide_num, slide in enumerate(prs.slides):
            slide_items = []

            # Extract text from shapes
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue

                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue

                    content_type = "heading" if para.level == 0 else "paragraph"
                    if text.endswith(":") and len(text) < 100:
                        content_type = "heading"

                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="pptx",
                        page_or_slide=slide_num + 1,
                        content_type=content_type,
                        content=text,
                        metadata={
                            "shape_type": shape.shape_type,
                            "position": {"left": shape.left, "top": shape.top}
                        }
                    ))

                # Check for speaker notes
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        items.append(ExtractedItem(
                            source_file=rel_path,
                            source_type="pptx",
                            page_or_slide=slide_num + 1,
                            content_type="speaker_notes",
                            content=notes,
                            metadata={}
                        ))

            # Extract tables
            for shape in slide.shapes:
                if shape.has_table:
                    table_data = []
                    for row in shape.table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        table_data.append(row_data)
                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="pptx",
                        page_or_slide=slide_num + 1,
                        content_type="table",
                        content=json.dumps(table_data),
                        metadata={}
                    ))

    except Exception as e:
        print(f"Error extracting PPTX {filepath}: {e}", file=sys.stderr)

    return items

def extract_docx(filepath: Path, rel_path: str) -> List[ExtractedItem]:
    """Extract content from DOCX."""
    items = []
    try:
        from docx import Document
        from docx.table import Table

        doc = Document(filepath)

        for para_num, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue

            content_type = "heading" if para.style.name.startswith("Heading") else "paragraph"

            items.append(ExtractedItem(
                source_file=rel_path,
                source_type="docx",
                page_or_slide=para_num + 1,  # Approximate
                content_type=content_type,
                content=text,
                metadata={"style": para.style.name}
            ))

        # Extract tables
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)
            items.append(ExtractedItem(
                source_file=rel_path,
                source_type="docx",
                page_or_slide=0,
                content_type="table",
                content=json.dumps(table_data),
                metadata={}
            ))

    except Exception as e:
        print(f"Error extracting DOCX {filepath}: {e}", file=sys.stderr)

    return items

def extract_text(filepath: Path, rel_path: str) -> List[ExtractedItem]:
    """Extract content from plain text/markdown files."""
    items = []
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()

        # Split by headings for markdown
        if filepath.suffix.lower() == '.md':
            sections = re.split(r'\n(?=#{1,6}\s+)', content)
            for sec_num, section in enumerate(sections):
                if not section.strip():
                    continue
                lines = section.strip().split('\n')
                heading = lines[0] if lines else ""
                body = '\n'.join(lines[1:]) if len(lines) > 1 else ""

                if heading.startswith('#'):
                    level = len(heading) - len(heading.lstrip('#'))
                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="md",
                        page_or_slide=sec_num + 1,
                        content_type="heading",
                        content=heading.lstrip('#').strip(),
                        metadata={"level": level}
                    ))

                if body:
                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="md",
                        page_or_slide=sec_num + 1,
                        content_type="paragraph",
                        content=body,
                        metadata={}
                    ))
        else:
            # Plain text - split by double newlines
            paragraphs = content.split('\n\n')
            for i, para in enumerate(paragraphs):
                para = para.strip()
                if para:
                    items.append(ExtractedItem(
                        source_file=rel_path,
                        source_type="txt",
                        page_or_slide=i + 1,
                        content_type="paragraph",
                        content=para,
                        metadata={}
                    ))

    except Exception as e:
        print(f"Error extracting text {filepath}: {e}", file=sys.stderr)

    return items

def classify_text_block(text: str, block: Dict) -> str:
    """Classify a text block by content type."""
    # Check for headings (large font, short, maybe all caps)
    if "spans" in block.get("lines", [{}])[0]:
        spans = block["lines"][0]["spans"]
        if spans:
            avg_size = sum(s.get("size", 12) for s in spans) / len(spans)
            if avg_size > 14 and len(text) < 120:
                return "heading"

    # Check for formulas (contains math-like patterns)
    if re.search(r'[∑∏∫∂∇∈∉⊂⊃∪∩∧∨¬∀∃∞≈≠≤≥±×÷]', text):
        return "formula"
    if re.search(r'\$[^$]+\$|\\\(|\\\[|\\begin\{', text):
        return "formula"

    # Check for code-like content
    if re.search(r'^\s*(def|function|class|if|for|while|return|import|#include)', text, re.MULTILINE):
        return "code"

    # Check for list items
    if re.match(r'^\s*[-•*]\s+', text) or re.match(r'^\s*\d+\.\s+', text):
        return "list"

    return "paragraph"

def extract_font_info(block: Dict) -> Dict:
    """Extract font information from a text block."""
    fonts = set()
    sizes = []
    for line in block.get("lines", []):
        for span in line.get("spans", []):
            fonts.add(span.get("font", ""))
            sizes.append(span.get("size", 12))
    return {
        "fonts": list(fonts),
        "avg_size": sum(sizes) / len(sizes) if sizes else 12,
        "max_size": max(sizes) if sizes else 12
    }

def extract_all(input_dir: str, work_dir: str) -> Dict[str, Any]:
    """Extract content from all files in inventory."""
    root = Path(input_dir).resolve()
    work_path = Path(work_dir).resolve()

    # Load inventory
    inventory_file = work_path / "inventory.json"
    if not inventory_file.exists():
        raise ValueError(f"Inventory not found: {inventory_file}")

    with open(inventory_file) as f:
        inventory = json.load(f)

    all_items = []
    extraction_stats = {"total_items": 0, "by_source_type": {}, "by_content_type": {}}

    for file_info in inventory["files"]:
        filepath = root / file_info["path"]
        if not filepath.exists():
            continue

        rel_path = file_info["path"]
        ext = file_info["extension"].lower()

        items = []
        if ext == '.pdf':
            items = extract_pdf(filepath, rel_path)
        elif ext in ('.ppt', '.pptx'):
            items = extract_pptx(filepath, rel_path)
        elif ext in ('.doc', '.docx'):
            items = extract_docx(filepath, rel_path)
        elif ext in ('.txt', '.md'):
            items = extract_text(filepath, rel_path)
        else:
            print(f"Skipping unsupported format: {ext}")

        # Convert to dict
        for item in items:
            all_items.append(asdict(item))
            extraction_stats["by_source_type"][item.source_type] = extraction_stats["by_source_type"].get(item.source_type, 0) + 1
            extraction_stats["by_content_type"][item.content_type] = extraction_stats["by_content_type"].get(item.content_type, 0) + 1

    extraction_stats["total_items"] = len(all_items)

    result = {
        "extraction_timestamp": datetime.now().isoformat(),
        "source_directory": str(root),
        "stats": extraction_stats,
        "items": all_items
    }

    # Save
    work_path.mkdir(parents=True, exist_ok=True)
    output_file = work_path / "sources.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"Extraction complete: {len(all_items)} items from {extraction_stats['by_source_type']}")
    return result

def main():
    if len(sys.argv) < 3:
        print("Usage: python extract_materials.py <input_directory> <work_directory>")
        sys.exit(1)

    input_dir = sys.argv[1]
    work_dir = sys.argv[2]

    try:
        extract_all(input_dir, work_dir)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()